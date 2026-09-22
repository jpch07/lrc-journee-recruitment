"""Isolated, real-HTTP event load test (never accepts a production HTTP target).

PowerShell:
  $env:EVALDAY_LOADTEST_DATABASE_URL='sqlite:///tmp/evalday_loadtest_smoke.db'
  python scripts/load_test_event.py --seconds 60 --report outputs/loadtest-smoke.json

For a hosted database, create an EMPTY, dedicated PostgreSQL database whose name
starts with ``evalday_loadtest_``. Supply its URI in the same environment variable.
The harness starts its OWN loopback-only uvicorn process, populates only synthetic
data, and leaves the database intact for inspection. It never drops tables,
switches production settings, connects to R2, or reads real accounts.

Default workload: 20 participants, 30 evaluator HTTP sessions, two admins, and
one restricted attendance operator. All five activities are open concurrently
(deliberately heavier than normal event operation), with 200 assigned tasks.
Real authentication, 5-second polling, drafts, final/idempotent submissions,
photo/cache access, general assessment autosaving, attendance writes, results,
and export are checked. A short run is NOT an endurance or Render-capacity test.
"""
from __future__ import annotations

import argparse
import asyncio
import base64
from collections import defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import time
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
PREFIX = "evalday_loadtest_"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1sAAAAASUVORK5CYII="
)


def validate_test_database(url: str, production_url: str = "") -> str:
    """Require an unmistakably test-only DB, and reject known production hosts."""
    if not url or any(value in url.lower() for value in ("cockroachlabs.cloud", "neon.tech", "onrender.com")):
        raise ValueError("An explicit isolated test database is required; production providers are blocked.")
    if url.startswith("sqlite:///"):
        filename = url[len("sqlite:///"):]
        path = Path(filename).resolve()
        if not path.stem.startswith(PREFIX) or path.exists():
            raise ValueError("SQLite requires a NEW file named evalday_loadtest_*.db.")
        if not path.parent.is_dir():
            raise ValueError("The SQLite parent directory must already exist.")
        return "sqlite:///" + path.as_posix()
    parsed = urlsplit(url)
    if parsed.scheme not in {"postgres", "postgresql", "postgresql+psycopg"}:
        raise ValueError("Only PostgreSQL and fresh SQLite test databases are supported.")
    if not parsed.hostname or not parsed.path.removeprefix("/").startswith(PREFIX):
        raise ValueError("PostgreSQL database name must start with evalday_loadtest_.")
    if production_url and parsed.hostname == urlsplit(production_url).hostname:
        if parsed.path == urlsplit(production_url).path:
            raise ValueError("The test database must not equal the production database.")
    return url.replace("postgres://", "postgresql+psycopg://", 1).replace(
        "postgresql://", "postgresql+psycopg://", 1
    )


def percentile(values: list[float], percent: int) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    return round(ordered[max(0, math.ceil(len(ordered) * percent / 100) - 1)], 2)


class Metrics:
    def __init__(self):
        self.routes = defaultdict(lambda: {"latencies": [], "status": defaultdict(int)})
        self.errors: list[str] = []

    def record(self, label: str, elapsed: float, status: int):
        row = self.routes[label]
        row["latencies"].append(elapsed * 1000)
        row["status"][str(status)] += 1

    def summary(self) -> dict:
        return {
            label: {
                "requests": len(row["latencies"]), "status": dict(row["status"]),
                "p50_ms": percentile(row["latencies"], 50),
                "p95_ms": percentile(row["latencies"], 95),
                "p99_ms": percentile(row["latencies"], 99),
                "max_ms": round(max(row["latencies"], default=0), 2),
            } for label, row in sorted(self.routes.items())
        }


def configure_isolated_environment(database_url: str):
    # Never inherit production object-storage, directory or messaging credentials.
    for key in list(os.environ):
        if key.startswith(("LRC_", "DATABASE_URL")):
            os.environ.pop(key)
    os.environ.update({
        "LRC_DATABASE_URL": database_url,
        "LRC_JOURNEE_ENV": "development",
        "LRC_JOURNEE_SESSION_SECRET": secrets.token_urlsafe(48),
        "LRC_JOURNEE_ADMIN_PASSWORD": secrets.token_urlsafe(24),
        "LRC_JOURNEE_COOKIE_SECURE": "false",
        "LRC_JOURNEE_TEST_TOOLS": "false",
        "LRC_RECRUIT_SHEET_ID": "",
        "LRC_JOURNEE_SESSION_HOURS": "48",
    })
    sys.path.insert(0, str(ROOT))


def seed_database(recruit_count: int, evaluator_count: int, password: str) -> dict:
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import inspect, select, text
    from app.db import Base, SessionLocal, engine
    from app import models as m
    from app.auth import hash_password
    from app.assessment_service import ensure_assessment_system
    from app.services import create_journey
    from app.tenant import select_system
    from app.rubric import ACTIVITY_ORDER, RUBRICS

    with engine.begin() as conn:
        schema = None if conn.dialect.name == "sqlite" else "journee_recruitment"
        existing = inspect(conn).get_table_names(schema=schema) if schema is None or schema in inspect(conn).get_schema_names() else []
        if existing:
            raise RuntimeError("Refusing to use a test database containing application tables.")
        if schema:
            conn.execute(text("CREATE SCHEMA IF NOT EXISTS journee_recruitment"))
        Base.metadata.create_all(conn)
        cfg = Config(str(ROOT / "alembic.ini"))
        cfg.set_main_option("script_location", str(ROOT / "migrations"))
        cfg.attributes["connection"] = conn
        command.stamp(cfg, "head")

    password_hash = hash_password(password)
    with SessionLocal() as db:
        system = ensure_assessment_system(db)
        system.name = "Synthetic Evalday Load Test"
        system.slug = "evalday-loadtest"
        select_system(system.id)
        journey = create_journey(db, "Synthetic 20/30 event", date.today(), 3, "Load test")
        journey.status = "active"
        journey.current_activity = "sport"
        recruits = [m.Recruit(
            journey_id=journey.id, name=f"Synthetic Recruit {i:02}", present=True,
            arrival_time=datetime.now(timezone.utc), photo_data=PNG, photo_type="image/png",
            photo_size=len(PNG), photo_sha256=hashlib.sha256(PNG).hexdigest(),
        ) for i in range(recruit_count)]
        db.add_all(recruits)
        evaluators = []
        usernames = []
        for i in range(evaluator_count):
            name = f"Synthetic Evaluator {i:02}"
            role = "overall" if i < evaluator_count // 2 else "dossard"
            directory = m.EvaluatorDirectory(system_id=system.id, name=name, default_role=role)
            db.add(directory)
            db.flush()
            evaluator = m.Evaluator(journey_id=journey.id, directory_id=directory.id, name=name,
                                    role=role, present=True)
            evaluators.append(evaluator)
            usernames.append(name)
            db.add_all([evaluator, m.UserAccount(
                system_id=system.id, username=name, directory_id=directory.id,
                password_hash=password_hash, evaluator_role=role, can_evaluate=True,
                must_change_password=False,
            )])
        for username in ["Synthetic Admin A", "Synthetic Admin B", "Synthetic Attendance"]:
            account = m.UserAccount(
                system_id=system.id, username=username, password_hash=password_hash,
                can_admin="Admin" in username, can_results="Admin" in username,
                can_evaluate=False, is_owner=username == "Synthetic Admin A", must_change_password=False,
            )
            db.add(account)
            db.flush()
            if username == "Synthetic Attendance":
                db.add(m.JourneyPermission(account_id=account.id, journey_id=journey.id, can_attendance=True))
        db.flush()
        expected = 0
        values = {}
        for code in ACTIVITY_ORDER:
            round_record = m.AssignmentRound(journey_id=journey.id, activity_code=code, version=1,
                status="published", seed="synthetic-loadtest", created_by="Load test")
            db.add(round_record)
            db.flush()
            state = db.scalar(select(m.ActivityState).where(
                m.ActivityState.journey_id == journey.id, m.ActivityState.code == code))
            state.status = "open"
            state.assignment_round_id = round_record.id
            for index in range(recruit_count * 2):
                db.add(m.Assignment(round_id=round_record.id, evaluator_id=evaluators[index % evaluator_count].id,
                    recruit_id=recruits[index % recruit_count].id, slot=index // recruit_count + 1,
                    room_number=index % 3 + 1))
                expected += 1
            rubric = RUBRICS[code]
            values[code] = {
                "responses": {c.key: 4 for c in rubric.criteria} if code != "sport" else {},
                "raw": {c.key: float(c.target) * 0.8 for c in rubric.criteria} if code == "sport" else {},
            }
        db.flush()
        attendance_token = db.get(m.RecruitAttendanceAccess, journey.id).token
        fixture = {"journey_id": journey.id, "recruit_ids": [r.id for r in recruits],
                   "evaluators": usernames, "expected_submissions": expected,
                   "attendance_token": attendance_token, "payloads": values}
        db.commit()
        fixture["database_bytes_before_workload"] = (
            db.scalar(text("SELECT pg_database_size(current_database())"))
            if engine.dialect.name != "sqlite" else None
        )
    engine.dispose()
    return fixture


class EventLoad:
    def __init__(self, base: str, fixture: dict, password: str, args):
        self.base, self.fixture, self.password, self.args = base, fixture, password, args
        self.metrics = Metrics()
        self.clients = []
        self.deadline = 0.0
        self.login_slots = asyncio.Semaphore(args.login_burst)

    async def progress(self):
        while True:
            await asyncio.sleep(60)
            statuses = defaultdict(int)
            for row in self.metrics.routes.values():
                for status, count in row["status"].items():
                    statuses[status] += count
            print(json.dumps({"progress": "running", "requests": sum(
                len(row["latencies"]) for row in self.metrics.routes.values()),
                "status_counts": dict(statuses),
                "failed_requests": sum(count for status, count in statuses.items()
                                       if status == "0" or int(status) >= 400),
                "steady_seconds_remaining": max(0, round(self.deadline - time.monotonic()))}), flush=True)

    async def request(self, client, method, path, label, *, expected=(200,), **kwargs):
        start = time.perf_counter()
        try:
            response = await client.request(method, path, **kwargs)
        except Exception as exc:
            self.metrics.record(label, time.perf_counter() - start, 0)
            # Deliberately omit exception strings, request bodies and URLs.
            raise RuntimeError(f"{label}: {type(exc).__name__}") from None
        self.metrics.record(label, time.perf_counter() - start, response.status_code)
        if response.status_code not in expected:
            raise RuntimeError(f"{label}: HTTP {response.status_code}")
        return response

    async def login(self, name):
        import httpx
        # Expire before uvicorn's five-second idle close. Otherwise exact five-
        # second polling can race a closing socket (browsers retry these GETs,
        # but HTTPX reports RemoteProtocolError instead).
        client = httpx.AsyncClient(base_url=self.base, timeout=self.args.timeout, trust_env=False,
            limits=httpx.Limits(max_connections=2, max_keepalive_connections=1, keepalive_expiry=2))
        self.clients.append(client)
        async with self.login_slots:
            response = await self.request(client, "POST", "/api/auth/login", "login", json={
                "username": name, "password": self.password, "recruitment": "evalday-loadtest"})
        client.headers["X-CSRF-Token"] = response.json()["csrfToken"]
        return client

    async def evaluator(self, client, index):
        home = (await self.request(client, "GET", "/api/evaluator/home", "evaluator.home")).json()
        tasks = [(a["code"], t) for a in home["activities"] for t in a["tasks"]]
        if not tasks:
            raise RuntimeError("Fixture evaluator has no assignments")
        photo = await self.request(client, "GET", tasks[0][1]["photoUrl"], "evaluator.photo")
        await self.request(client, "GET", tasks[0][1]["photoUrl"], "evaluator.photo.cached",
                           headers={"If-None-Match": photo.headers["etag"]}, expected=(304,))
        versions = {}
        iterations = 0
        while time.monotonic() < self.deadline:
            await self.request(client, "GET", "/api/evaluator/updates", "evaluator.poll")
            code, task = tasks[iterations % len(tasks)]
            task_id = task["assignmentId"]
            body = {**self.fixture["payloads"][code], "comments": f"Synthetic draft {iterations}",
                    "client_version": versions.get(task_id)}
            draft = await self.request(client, "PUT", f"/api/evaluator/tasks/{task_id}/draft",
                "evaluator.draft", json=body, headers={"Idempotency-Key": f"draft-{index}-{iterations}"})
            versions[task_id] = draft.json()["submission"]["version"]
            iterations += 1
            await asyncio.sleep(min(self.args.poll_seconds, max(0, self.deadline - time.monotonic())))
        # Finish all assigned tasks. Repeat the exact request to check idempotency.
        for code, task in tasks:
            task_id = task["assignmentId"]
            body = {**self.fixture["payloads"][code], "comments": "Synthetic final evaluation",
                    "client_version": versions.get(task_id)}
            headers = {"Idempotency-Key": f"final-{task_id}"}
            response = await self.request(client, "POST", f"/api/evaluator/tasks/{task_id}/submit",
                "evaluator.submit", json=body, headers=headers)
            duplicate = await self.request(client, "POST", f"/api/evaluator/tasks/{task_id}/submit",
                "evaluator.submit.retry", json=body, headers=headers)
            if response.json() != duplicate.json():
                raise RuntimeError("Duplicate submission did not return the same result")
            if response.json()["submission"]["score"] != 4:
                raise RuntimeError("Unexpected evaluation score")

    async def admin(self, client, index):
        base = f"/api/admin/journeys/{self.fixture['journey_id']}"
        iterations = 0
        while time.monotonic() < self.deadline:
            suffix = ["dashboard", "results", "monitoring/sport", "assignments/escape_room"][iterations % 4]
            await self.request(client, "GET", f"{base}/{suffix}", f"admin.{suffix.split('/')[0]}")
            recruit_id = self.fixture["recruit_ids"][(iterations * 2 + index) % len(self.fixture["recruit_ids"])]
            profile_path = f"{base}/recruits/{recruit_id}/profile"
            profile = (await self.request(client, "GET", profile_path, "admin.profile")).json()
            await self.request(client, "PUT", profile_path, "admin.assessment.autosave", json={
                "punctuality": 0.8, "respect": 0.8, "seriousness": 0.8,
                "comment": "Synthetic assessment", "notes": "Synthetic note",
                "base_version": profile["assessment"]["version"],
            })
            iterations += 1
            await asyncio.sleep(min(self.args.poll_seconds * 2, max(0, self.deadline - time.monotonic())))

    async def attendance(self, client):
        await self.request(client, "POST",
            f"/api/public/recruit-attendance/{self.fixture['attendance_token']}/select", "attendance.select")
        iteration = 0
        while time.monotonic() < self.deadline:
            roster = (await self.request(client, "GET", "/api/recruit-attendance/recruits", "attendance.poll")).json()
            recruit = roster["recruits"][iteration % len(roster["recruits"])]
            await self.request(client, "PATCH", f"/api/recruit-attendance/recruits/{recruit['id']}",
                "attendance.autosave", json={"base_version": recruit["version"],
                                            "attendance_comment": f"Synthetic arrival note {iteration}"})
            iteration += 1
            await asyncio.sleep(min(self.args.poll_seconds, max(0, self.deadline - time.monotonic())))

    async def run(self):
        reporter = None
        try:
            names = [*self.fixture["evaluators"], "Synthetic Admin A", "Synthetic Admin B", "Synthetic Attendance"]
            clients = await asyncio.gather(*(self.login(name) for name in names), return_exceptions=True)
            login_errors = [str(client) for client in clients if isinstance(client, Exception)]
            if login_errors:
                # Finish the whole login burst before closing clients; otherwise
                # one 429 would turn unrelated in-flight requests into fake I/O failures.
                self.metrics.errors.extend(login_errors)
                return
            self.deadline = time.monotonic() + self.args.seconds
            reporter = asyncio.create_task(self.progress())
            results = await asyncio.gather(
                *(self.evaluator(client, i) for i, client in enumerate(clients[:-3])),
                self.admin(clients[-3], 0), self.admin(clients[-2], 1), self.attendance(clients[-1]),
                return_exceptions=True,
            )
            self.metrics.errors.extend(str(r) for r in results if isinstance(r, Exception))
            # Normal lifecycle API closes stages before generating completed-event report.
            if not self.metrics.errors:
                prefix = f"/api/admin/journeys/{self.fixture['journey_id']}"
                results = (await self.request(clients[-3], "GET", prefix + "/results", "admin.final.results")).json()
                for row in results["rows"]:
                    for value in row["activities"].values():
                        if value["score"] != 4 or value["submitted"] != 2:
                            raise RuntimeError("Rankings do not match exactly two submitted scorecards per recruit/activity")
                for activity in self.fixture["payloads"]:
                    await self.request(clients[-3], "POST", prefix + f"/activities/{activity}/close",
                                       "admin.activity.close", json={"reason": "Synthetic load test complete"})
                journey = (await self.request(clients[-3], "GET", prefix, "admin.journey")).json()
                await self.request(clients[-3], "PATCH", prefix, "admin.journey.complete",
                                   json={"status": "completed", "base_version": journey["version"]})
                report = await self.request(clients[-3], "GET", prefix + "/results.xlsx", "admin.export")
                import io
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(report.content), read_only=True, data_only=False)
                try:
                    if workbook.sheetnames != ["Attendance", "Results", "Recruit Profiles"]:
                        raise RuntimeError("Report did not contain the three management sheets")
                    values = {value for row in workbook["Attendance"].iter_rows(values_only=True)
                              for value in row if isinstance(value, str)}
                    for i in range(self.args.recruits):
                        if f"Synthetic Recruit {i:02}" not in values:
                            raise RuntimeError("Completed-event report lost a recruit")
                finally:
                    workbook.close()
        except Exception as exc:
            self.metrics.errors.append(str(exc))
        finally:
            if reporter:
                reporter.cancel()
                await asyncio.gather(reporter, return_exceptions=True)
            await asyncio.gather(*(c.aclose() for c in self.clients))


def verify_database(fixture) -> dict:
    from sqlalchemy import func, select, text
    from app.db import SessionLocal, engine
    from app.models import EvaluationSubmission, SubmissionVersion
    with SessionLocal() as db:
        submitted = db.scalar(select(func.count()).select_from(EvaluationSubmission).where(
            EvaluationSubmission.status.in_(["submitted", "locked"])))
        wrong = db.scalar(select(func.count()).select_from(EvaluationSubmission).where(
            EvaluationSubmission.score != 4))
        duplicates = db.execute(select(EvaluationSubmission.assignment_id, func.count()).group_by(
            EvaluationSubmission.assignment_id).having(func.count() > 1)).all()
        history = db.scalar(select(func.count()).select_from(SubmissionVersion))
        size = db.scalar(text("SELECT pg_database_size(current_database())")) if engine.dialect.name != "sqlite" else None
    engine.dispose()
    return {"expected_submissions": fixture["expected_submissions"], "submitted": submitted,
            "wrong_scores": wrong, "duplicate_assignments": len(duplicates), "history_rows": history,
            "database_bytes_after": size,
            "passed": submitted == fixture["expected_submissions"] and wrong == 0 and not duplicates}


async def run_bounded_workload(load: EventLoad, deadline_seconds: float):
    try:
        await asyncio.wait_for(load.run(), timeout=deadline_seconds)
    except asyncio.TimeoutError:
        load.metrics.errors.append("Workload exceeded its bounded runtime; outstanding test requests cancelled")


def main() -> int:
    import httpx
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--seconds", type=float, default=60)
    parser.add_argument("--poll-seconds", type=float, default=5)
    parser.add_argument("--timeout", type=float, default=60)
    parser.add_argument("--grace-seconds", type=float, default=300,
                        help="Maximum extra time for HTTP login, final submission, verification/export beyond --seconds.")
    parser.add_argument("--login-burst", type=int, default=3)
    parser.add_argument("--recruits", type=int, default=20)
    parser.add_argument("--evaluators", type=int, default=30)
    parser.add_argument("--report", type=Path, default=ROOT / "outputs" / "loadtest-event.json")
    args = parser.parse_args()
    if args.seconds < 1 or args.poll_seconds < 1 or args.login_burst < 1 or args.grace_seconds < 1:
        parser.error("Durations and login concurrency must be positive.")
    if not 2 <= args.evaluators <= args.recruits * 2 or args.evaluators <= args.recruits:
        parser.error("This fixture needs recruits < evaluators <= 2*recruits.")
    database_url = validate_test_database(os.getenv("EVALDAY_LOADTEST_DATABASE_URL", ""), os.getenv("LRC_DATABASE_URL", ""))
    configure_isolated_environment(database_url)
    password = secrets.token_urlsafe(24)
    fixture = seed_database(args.recruits, args.evaluators, password)
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    base = f"http://127.0.0.1:{port}"
    # No raw application logs: traceback locals or connection exceptions can contain credentials.
    process = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1",
        "--port", str(port), "--no-access-log", "--log-level", "critical"], cwd=ROOT,
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    started = time.monotonic()
    started_utc = datetime.now(timezone.utc).isoformat()
    try:
        ready = False
        for _ in range(120):
            if process.poll() is not None:
                raise RuntimeError("Isolated application process exited before becoming ready")
            try:
                ready = httpx.get(base + "/health/ready", timeout=2, trust_env=False).status_code == 200
            except httpx.HTTPError:
                pass
            if ready:
                break
            time.sleep(0.5)
        if not ready:
            raise RuntimeError("Isolated application startup exceeded readiness timeout")
        load = EventLoad(base, fixture, password, args)
        asyncio.run(run_bounded_workload(load, args.seconds + args.grace_seconds))
        verification = verify_database(fixture)
        metrics = load.metrics.summary()
        report = {
            "started_utc": started_utc,
            "elapsed_seconds": round(time.monotonic() - started, 2),
            "steady_workload_seconds": args.seconds,
            "maximum_http_workload_seconds": args.seconds + args.grace_seconds,
            "recruits": args.recruits, "concurrent_evaluators": args.evaluators,
            "admin_sessions": 2, "attendance_sessions": 1, "poll_interval_seconds": args.poll_seconds,
            "login_concurrency": args.login_burst,
            "database_kind": "sqlite" if database_url.startswith("sqlite") else "hosted_postgresql",
            "http_host": "isolated_loopback_process", "verification": verification,
            "database_bytes_before_workload": fixture.get("database_bytes_before_workload"),
            "routes": metrics, "errors": load.metrics.errors,
            "passed": verification["passed"] and not load.metrics.errors,
            "scope_limits": ["Does not measure Render CPU/RAM, internet clients, or regional server-to-DB latency.",
                "Synthetic tiny photos use the legacy DB photo endpoint, not production R2.",
                "All five stages intentionally open together; does not validate assignment-generation rules.",
                "Short runs cannot establish 12-hour endurance or guarantee availability."],
        }
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(json.dumps(report, indent=2))
        return 0 if report["passed"] else 1
    finally:
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=5)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValueError, RuntimeError) as error:
        print(f"Load test stopped: {error}", file=sys.stderr)
        raise SystemExit(2)
    except Exception as error:
        # Never let a driver exception expose a credential-bearing connection URI.
        print(f"Load test stopped during setup/verification: {type(error).__name__}", file=sys.stderr)
        raise SystemExit(2)
